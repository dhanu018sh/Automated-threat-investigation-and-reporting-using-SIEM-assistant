import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

pptx_path = r"D:\antigravity\scratch\siem-assistant\Automated_Threat_Investigation_SIEM_Assistant_Phase_2_Presentation.pptx"
prs = pptx.Presentation(pptx_path)

s16 = prs.slides[15] # Slide 16

NAVY_BLUE = RGBColor(16, 32, 64)
TEXT_DARK = RGBColor(30, 30, 30)

# Clear Rectangle 7
for sh in s16.shapes:
    if sh.name == "Rectangle 7" and sh.has_text_frame:
        sh.text_frame.clear()
        print("Cleared text from Rectangle 7.")

# Find TextBox 9
tb9 = None
for sh in s16.shapes:
    if sh.name == "TextBox 9":
        tb9 = sh
        break

if tb9:
    tb9.left = Inches(0.8)
    tb9.top = Inches(1.4)
    tb9.width = Inches(11.8)
    tb9.height = Inches(5.2)
    
    tf = tb9.text_frame
    tf.word_wrap = True
    tf.clear()
    
    outcomes = [
        ("1. Mean Time to Detect (MTTD) Reduction", "Threat triage turnaround reduced from 45 minutes of manual log searching to under 3 seconds using real-time Random Forest ML evaluation on incoming SIEM event streams."),
        ("2. False Positive Reduction", "Achieves 60%–70% reduction in false positive alerts through 100-tree ensemble consensus voting, overcoming single-threshold rule triggers."),
        ("3. Analyst Productivity Improvement", "Tier-1 analysts investigate complex security incidents entirely in plain English without requiring SPL, KQL, or AQL syntax expertise, reducing turnaround time by 40%."),
        ("4. Automated Incident Documentation", "Instant generation of structured incident reports (Plain Text .txt, Markdown, HTML, PDF) with extracted IOCs, MITRE ATT&CK technique IDs, and recommended remediation actions."),
        ("5. Multi-Platform SIEM Coverage", "Pluggable SIEM connector abstraction supporting live Windows Event Logs, Splunk, Elastic Stack, Microsoft Sentinel, IBM QRadar, and RESTful log sources."),
        ("6. Persistent SOC Knowledge Base", "Investigation threads and generated threat reports are persistently archived in SQLite/localStorage, creating a searchable and auditable incident knowledge base over time.")
    ]
    
    for idx, (title, desc) in enumerate(outcomes):
        p_title = tf.add_paragraph() if idx > 0 else tf.paragraphs[0]
        p_title.text = title
        p_title.font.name = "Calibri"
        p_title.font.size = Pt(12.5)
        p_title.font.bold = True
        p_title.font.color.rgb = NAVY_BLUE
        p_title.space_before = Pt(4)
        
        p_desc = tf.add_paragraph()
        p_desc.text = desc
        p_desc.font.name = "Calibri"
        p_desc.font.size = Pt(10.5)
        p_desc.font.color.rgb = TEXT_DARK
        p_desc.space_after = Pt(4)
    
    print("Populated clean text in TextBox 9.")

prs.save(pptx_path)
print("Saved corrected presentation to:", pptx_path)
