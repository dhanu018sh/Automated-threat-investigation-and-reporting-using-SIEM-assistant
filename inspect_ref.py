import os
import sys
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

ref_path = r"C:\Users\Dhanush\.gemini\antigravity\brain\f8dd61c8-d9ed-419f-a1e2-6cfa12754ef2\.user_uploaded\media_1791270872056.pptx"
prs = pptx.Presentation(ref_path)

print(f"Total slides in reference: {len(prs.slides)}")
for idx, s in enumerate(prs.slides):
    t = s.shapes.title.text.strip().replace('\n', ' ') if s.shapes.title else "No Title"
    print(f"Slide {idx+1}: {t}")
