import pptx

ref_path = r"C:\Users\Dhanush\.gemini\antigravity\brain\f8dd61c8-d9ed-419f-a1e2-6cfa12754ef2\.user_uploaded\media_1791270872056.pptx"
prs = pptx.Presentation(ref_path)

for idx in [11, 12, 13, 14, 15]:
    slide = prs.slides[idx]
    print(f"\n==================== SLIDE {idx+1} ====================")
    for s in slide.shapes:
        t = s.text.strip().replace('\n', ' ') if s.has_text_frame else ''
        print(f"Shape: id={s.shape_id}, name={s.name}, type={s.shape_type}, pos=({s.left},{s.top},{s.width},{s.height})")
        if t:
            print(f"   Text: {t[:80]}")
